from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
import re

ROOT = Path(r'C:/Users/mirav/OneDrive/Desktop/MiraverseOSx')
SOURCE = Path(r'C:/Users/mirav/Downloads/MiraverseOSx_Expanded_Application_Encyclopedia_2026.docx')
OUT = ROOT / 'docs/MiraverseOSx_World_and_Application_Bible_2026_Updated.docx'
doc = Document(SOURCE)
p = list(doc.paragraphs)

# Update the existing encyclopedia in place, retaining its historical chapters.
replacements = {
1: 'Meridion World Lore and Application Bible',
2: 'World history and the connected desktop experience',
3: 'October 23 2026  |  Edition 1.1',
6: 'Meridion is described here as a world in its own right. This volume combines a historical chronicle with an encyclopedia of the services through which citizens live, work, communicate, and investigate. It distinguishes established lore, institutional interpretation, unresolved mysteries, and the current playable implementation. Chapters XXIV and XXV record the connected desktop as of October 23, 2026.',
8: 'Contents',
128: 'The historical chronicle continues into the living world and its institutions below. Unresolved origins remain open; interface changes do not establish new historical answers.',
130: 'Part Two The Living World',
154: 'The travel design allows direct movement and optional faster passage between eligible destinations. Current daily play centers on desktop navigation, location records, and contextual exploration. A legacy field-scene prototype exists, but it does not establish a fully traversable city or a requirement to use playable scenes for the connected investigation.',
163: 'A new local citizen begins with one personal AureMail mailbox. Institutional sending requires both an active organization membership and an explicitly issued institutional account. Comms supplies authorized organization channels; phone Messages holds personal conversations and drafts. Reading a notice does not confer membership or issue an address.',
171: 'The wider information network. Citizens visit fictional publications, schools, banks, businesses, and public resources. VerseNet also opens the shared Service Directory and an optional external website viewer. Real websites remain separate from Meridion records and lore.',
177: 'Employment and civic opportunities. WORKNET owns careers, accepted work orders, civic jobs, and civic notices. Former notice-board entry points lead here. Professional duties require their own evidence and completion records; viewing a listing alone does not certify work.',
187: 'The former name of the personal profile now consolidated into Citizen Hub. MIRROR links resolve to that same personal window, where the citizen can view their own profile, skills, and eligible records. It is not a second identity authority or a separate population editor.',
203: 'Personal and official documents. VAULT retains source records, assignment files, mail attachments, and saved materials. Notes, logs, maps, forms, clinical records, archives, technical files, and corrupted materials have distinct readers. A saved attachment retains its source content rather than becoming a generic civic form.',
205: 'Everyday personal communication. Phone Messages stores personal threads and per-contact drafts across closing and reloading. Earlier Comms direct conversations are imported once, while Comms retains institutional discussion and links to the phone.',
218: 'Daily life need not become an assignment. HOMECRAFT currently supports manually chosen cooking, tidying, rearranging, social time, and browsing furnishings, alongside its existing rest and purchase systems. Familiarity can simplify repeated routines and routine time advances simulation segments. These ordinary routines do not issue work rewards or XP. Fuller domestic and social simulation remains open for expansion.',
228: 'The application encyclopedia and dated implementation chapters complete this volume. They explain how current services express the world without assigning solutions to its unanswered historical questions.',
230: 'XXI The Digital Life of Meridion',
232: 'Entries distinguish a service owner from sections and legacy names that lead to it. Current behavior is described where implemented. Institutional boundaries express the world design; the local administrator tools are simulated conveniences and do not constitute production authentication or complete permission enforcement.',
248: 'Inbox and folders; contacts and threads; formatted composer and drafts; attachment reader; rules and preferences; Accounts and memberships dialog.',
250: 'Read, compose, reply, organize, archive, recover, and save correspondence. Tiptap supports formatted composition. Assignment confirmations and completion notices share stable work-order references. The Broad Street case supplies Mateo Cruz\'s requested clarification as one reusable correspondence thread.',
252: 'New saves begin with citizen@auremail.aure. Institutional sending requires an active matching membership and separately issued account. Suspended or revoked membership blocks sending; a suspended account also blocks sending. Older messages and account history remain readable, but historic account notices do not prove current authority.',
254: 'VAULT preserves an attachment\'s actual content and provenance. WORKNET issues deduplicated work correspondence. Account issuance and organization membership are managed separately; these are fictional addresses rather than real internet email delivery.',
259: 'Public broadcasts; institution and community channels; workplace coordination; organization access context; Accounts and memberships dialog; links to personal phone threads.',
261: 'Read public channels and participate in organization conversations when authorized. Restricted messages are hidden without membership, and posting rechecks access. Existing corruption diagnostics require genuine corruption and access before awarding their existing compensation.',
263: 'Accepted employment contracts establish the matching workplace membership unless an explicit suspension or revocation overrides it. Rank, reputation, an affiliation label, or administrator status alone does not establish membership. Cyacademy membership can be issued administratively; a complete admission process and assignment-specific channel compartments remain future work.',
270: 'Persistent personal messages and contacts; per-contact drafts; quick replies and dialogue choices; phone preferences and notifications.',
272: 'Continue personal threads after closing the phone or reloading. Earlier Comms direct messages migrate once with stable imported IDs, without deleting their original history. Text lost from the former memory-only phone before this change cannot be reconstructed.',
281: 'Address bar and tabs; in-world portals and public search; bookmarks and history; shared Service Directory; restricted external website viewer.',
283: 'Visit fictional services, follow case-aware source links, search public publications, and open external services within the workspace or separately. The fixed First Federal maintenance notice can be cited into Broad Street. Private evidence and answer rules do not enter public search.',
285: 'Institutional destinations retain their own access checks. External sites may block embedded viewing and require a separate browser tab or window. External pages receive no automatic citizen or case payload. A downloaded HTML tool can require supporting assets or a local server; its sandboxed preview is not a native application installation.',
287: 'The Service Directory collection is shared with DGAOS. Eleven default entries cover medical, police, forensic, legal, and fictional-data tools; MedCore awaits a configured link. Links can be searched, filtered, edited, added, or removed, and local HTML files can be selected. The standalone HTML export remains an optional separate artifact, not the integrated desktop.',
314: 'Career and qualification records; employment contracts; work orders and staged tasks; civic jobs and notices; formatted report drafts; review feedback; accepted submissions and compensation receipts.',
316: 'Accept eligible work, acknowledge a briefing, perform assignment-specific stages, and retain verified outputs. Broad Street adds a six-source investigation, structured findings with citations, a formatted explanation, revision feedback, and once-only finalization. Reading or saving materials does not itself pay compensation.',
318: 'Employment, institutional clearance, and assignment conditions retain separate meanings. Broad Street requires existing DGA or TRACE access at clearance 2. Promotion checks relevant skills and demonstrated experience alongside applicable qualifications, reputation, and approvals. An accepted assignment does not issue institutional credentials.',
320: 'AureMail handles correspondence, VAULT holds original assignment documents, DGAOS owns evidence, and VerseNet supplies public pages. WORKNET owns report review and compensation. Stable work-order and source references connect them without copying a base reward into each app.',
338: 'Arrange furnishings, buy eligible objects at their recorded costs, rest, and choose daily routines. Ordinary cooking, tidying, rearranging, and social time have their own history and familiarity; they do not create assignments or repeat existing upgrade rewards.',
347: 'Personal profile; civic service links; read-only linked AURA identity; personal document issuance history; inventory, skills, associations, and residence context.',
349: 'Review the current player\'s own linked record and issuance metadata, then open the owning service for housing, healthcare, work, archives, or education. An unlinked profile remains explicit; the hub does not fall back to another citizen selected in an agency app.',
351: 'AURA linkage uses an existing stable person identifier, not a name match. Browsing another registry subject does not relink the player. A local administrator can explicitly link an eligible existing identity. Calibration is a gameplay profile step, not official identity issuance, and does not replace existing identifiers or repay setup rewards.',
353: 'Citizen Hub consolidates the former MIRROR and CIVINET entry points. Legacy aliases and the old CIVINET web routes open the same personal workspace. AURA remains the separate population authority. Civic support retains its original one-time claim rather than becoming a recurring stipend.',
365: '13 MIRROR Legacy Entry',
366: 'Compatibility with the consolidated Citizen Hub',
367: 'MIRROR is retained as a legacy name for the personal workspace now presented as Citizen Hub. It is not an additional current application with a competing identity record.',
369: 'The existing personal profile, development, appearance, and authorized self-record within Citizen Hub.',
371: 'Use older MIRROR links or shortcuts to reach the same Citizen Hub window. Existing personal progress remains available.',
373: 'Personal identity remains separate from the shared active agency subject. MIRROR compatibility does not authorize editing another citizen.',
375: 'Citizen Hub owns the personal experience; AURA owns the canonical population record. The existing internal passport identifier preserves saved links and integrations.',
382: 'Inspect and produce supported official document types from canonical records. Personal issuance history preserves its original name, issuer, reference, timestamp, and hash; a historical metadata log is not presented as a recoverable original rendered file.',
391: 'Command Core; case queues and dossiers; Broad Street evidence notebook; agency phone, events, and feed; embedded TRACE; Directory toolbar button beside PHONE and FEED.',
393: 'Develop authorized cases, reopen original evidence, acquire case sources, review professional records, and use the embedded laboratory. Directory opens the existing shared service collection and returns from TRACE to Command Core. Its extra category column is omitted inside DGAOS so it does not consume the case workspace.',
395: 'Case-specific checks still apply to access and review. The signed-in local world offers a full-access credential action that enables administrator status, clearance 5, DGA verification, and active listed institution memberships without replacing character progress. It does not separately issue every mail account or turn administrator convenience into established world law.',
397: 'TRACE remains inside DGAOS. Legacy TRACE launch links focus the owning agency window. AURA, Document Studio, and MEDSTAT share canonical citizen updates; case links preserve the selected investigation and source.',
402: 'Case-linked forensic stations; evidence readers; autosaved examination notes; comparisons and report drafts; formal findings transmission; restored case and station context.',
404: 'Examine evidence, retain drafts, compare observations, and explicitly transmit formal findings. TRACE\'s citizen registry now awaits canonical records, refreshes on shared citizen updates, and opens the current AURA and Document Studio destinations.',
406: 'Draft progress and formally transmitted findings remain distinct. Broad Street uses dedicated authored source and report views; its report acceptance does not import the older courtroom preset as a verdict.',
424: 'Healthcare portal; patient records and clinical care; DGAOS MEDSTAT with ECG and chart views; canonical MAID medical links and summaries; profession-specific work records.',
426: 'Seek personal care through the owning service or perform supported professional duties. MEDSTAT loads citizen records asynchronously and follows shared citizen updates. Faith Medical and WORKNET share the Mercer care receipt so recording the clinical assignment does not pay again for the same treatment.',
446: 'Historical entries, people, places, factions, and chronology; Aureline exploration; Discovery Journal; contextual observations and leads.',
448: 'Study history and explore public Aureline locations through the desktop. Observations, unexplained objects, optional quest offers, adventure invitations, and location keepsakes have distinct presentation. Some discoveries require prior knowledge or an established connection.',
457: 'Home and quick-access folders; assignment case folders; saved mail and downloaded records; distinct note, log, map, form, clinical, archive, technical, and alert readers; document security and provenance.',
459: 'Reopen original materials, retain faithful attachment copies, inspect maps and structured records, and explicitly acquire eligible case sources. Document action buttons lead to the owning civic, medical, campus, or system workspace instead of silently issuing credentials or repeat credits.',
461: 'Encrypted content records an inspection only after it becomes readable through its supported access path. A driver file reference alone does not reveal protected lineage history. Having a copy does not override the case\'s authorization or make an unrelated source a valid citation.',
463: 'AureMail and VAULT share attachment content. WORKNET supplies briefings, checklists, and receipts; DGAOS evidence and public VerseNet records retain their case links. Completed rewards and discoveries survive reopening.',
468: 'Desktop appearance; credential controls for the signed-in local world; system monitor and console; Save and sync status, manual saving, and conflict recovery.',
470: 'Choose one of five saved color collections, manage shortcuts and windows, inspect system information, and review local database synchronization. Normal app windows use a 16 to 9 fit; maximizing fills the available desktop, and restoring returns to normal geometry.',
472: 'The full-access option is a local simulated administrator convenience. Current access checks and membership rules still have individual scopes, and complete production role enforcement remains unfinished. Local world snapshots include only listed game records, excluding unrelated storage and credentials.',
474: 'The desktop fills the browser edge to edge. Shared styles remove unintended native beveled button borders while preserving intentional boundaries and keyboard focus. App-level error recovery keeps an individual rendering failure from replacing the whole desktop.',
485: 'Creator journal state participates in the existing local world snapshot. This persistence support does not settle pause, resume, publishing, or authority to modify canon.',
488: 'AURA retains the authoritative identity. Citizen Hub presents the current player\'s own eligible services and read-only linked record; MIRROR is its legacy name. Primary and Secondary Houses retain their established meanings. No interface alone bypasses the House Tribunal.',
492: 'A resident accepts eligible work in WORKNET, acknowledges its briefing, and receives related correspondence and files. In Broad Street, VAULT records, DGAOS laboratory findings, an AureMail witness clarification, and a public VerseNet notice support the report returned to WORKNET. Review can request revision before the single final payment.',
500: 'Pulse\'s independent purpose, Creator Studio\'s effect on the living world, the banking model, full agency permissions, deeper professional and domestic simulation, and the Tribunal\'s detailed procedures remain open. Current mailbox issuance, persistent phone Messages, Citizen Hub consolidation, and Broad Street are implemented and no longer described as unspecified proposals.',
501: 'Shared addresses, varied residence detail, odd-and-even numbering, and realm-specific conventions continue to guide geography. Runtime postal prefixes and prototype location records do not by themselves settle every canonical designation. Current play emphasizes the connected desktop; a complete traversable city and further playable scenes are outside the present investigation scope.',
}
for index, text in replacements.items():
    p[index].text = text

# A former end-of-volume sentence otherwise spills alone ahead of the next part.
p[228]._element.getparent().remove(p[228]._element)

# Extend the existing contents without discarding its historical entries.
for text in [
    'XIII Addresses cities and the geography of everyday life',
    'XIV Housing property and domestic life',
    'XV Movement travel and discovery',
    'XVI MiraverseOSx and its digital life',
    'XVII Careers education and playable responsibilities',
    'XVIII Community relationships and institutions',
    'XIX The economy and public life',
    'XX Boundaries of established knowledge',
    'XXI The digital life of Meridion application encyclopedia',
    'XXII How a citizen moves through the services',
    'XXIII The open edges of everyday life',
    'XXIV Connected desktop implementation October 23 2026',
    'XXV Development record and verification',
]:
    p[22].insert_paragraph_before(text, style='Normal')

def heading(text, level=2):
    return doc.add_heading(text, level)

def para(text):
    return doc.add_paragraph(text)

def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers))
    for c,s in zip(t.rows[0].cells,headers): c.text=s
    for row in rows:
        for c,s in zip(t.add_row().cells,row): c.text=s
    if widths:
        from docx.shared import Inches
        t.autofit=False
        for row in t.rows:
            for cell,w in zip(row.cells,widths): cell.width=Inches(w)
        for col,w in zip(t.columns,widths): col.width=Inches(w)
    return t

doc.add_page_break()
heading('XXIV Connected Desktop Implementation October 23 2026',1)
para('This chapter records the current local player experience. The historical chapters remain the lore reference; this dated chapter describes the implemented mechanics and their limits. Completed examples demonstrate the connected desktop without establishing a final campaign, a complete city, or new answers to the world\'s mysteries.')

heading('One owner for each service')
table(['Former entry or section','Current owner'],[
('TRACE standalone launch','DGAOS with embedded TRACE laboratory'),
('CHRONICLE and QUESTNOTICE','WORKNET civic notices and civic jobs'),
('SpellForge','Cyacademy Magic practice'),
('MIRROR and CIVINET','Citizen Hub with the existing personal workspace'),
('Personal Comms conversations','Phone Messages with retained conversation history'),
('Service Directory','Shared VerseNet and DGAOS collection'),
('Pulse social content','MAI.SPACE; Pulse retains a status entry'),
], [2.45,4.05])
para('Legacy app identifiers and saved shortcuts lead to the owning workspace rather than creating duplicate applications. Existing records, claims, and unfinished work remain intact. The standalone service HTML source is retained because its export builder still uses it.')

heading('Broad Street connected investigation')
para('Broad Street Vault Breach uses CASE-2026-002, file reference FS-2026-00422, and WORKNET assignment ms-case-002. Its original people, chronology, and evidence IDs are retained. The case\'s 2026 record references are preserved as authored case data; they do not redefine Meridion\'s founding calendar. The target session is 20 to 30 minutes, without a deadline; this duration has not been established by a player study.')
para('An eligible player opens WORKNET Global Affairs, accepts the assignment, and acknowledges the briefing. Existing DGA or TRACE access and clearance 2 are required. A deduplicated AureMail opportunity dispatch can point to the work. Accepting does not issue credentials. The resulting VAULT folder contains a briefing, checklist, access audit, and thermal log.')
table(['Owning app','Source and collection action'],[
('VAULT','Access audit; open the original and acquire it.'),
('VAULT','Facilities thermal log; open the original and acquire it.'),
('DGAOS','EV-0010 metallurgy findings; acquire from the case source reader.'),
('DGAOS','EV-0011 DNA findings; acquire from the case source reader.'),
('AureMail','Request Mateo Cruz\'s clarification, read the thread, and acquire the statement.'),
('VerseNet','Search First Federal maintenance; open the public notice and cite it into the case.'),
], [1.15,5.35])
para('The evidence notebook tracks acquired sources, remaining objectives, original-source links, and optional hints. Sources can be examined in different orders. Repeated witness requests reopen the same thread. Reading, downloading, or saving a source does not automatically acquire it or produce a payout.')
para('The fixed public notice is handled internally at https://firstfederal.aure/maintenance. Only its public publication enters general search. Private sources, solution rules, and protected case records stay outside the general index. The investigation requires no real internet research.')

heading('Report review and compensation')
para('WORKNET supplies a Tiptap report editor and three structured finding questions. Each finding has its own selected supporting citations. Review checks known and acquired source IDs, case relevance, required evidence relationships, the chosen findings, and at least 80 characters of player-written explanation. Generated headings and placeholders do not count. It does not use AI or keyword matching to judge the prose.')
para('The review requires care about attribution: a replayed credential is not by itself proof that its registered holder participated, and an association is not a complete reconstruction of the crime. Feedback identifies missing or unsupported findings and citations without writing the answer. Revisions are unlimited and carry no reward or rating penalty. Acceptance records a supported investigation, not a criminal verdict.')
para('An accepted current draft enables normal WORKNET finalization for 300 credits, 150 XP, and zero bits, with the existing 10-stamina cost. Payment occurs once. Reading, requests, downloads, accepted reviews, and reopened reports issue no additional payment. Older paid or closed versions cannot become fresh paid work.')
para('Formatted drafts and submitted versions are retained. Editing an accepted draft invalidates its current acceptance receipt and requires review again. Unchanged repeated submissions reuse their earlier review. Restoring an earlier accepted draft and resubmitting can restore acceptance without duplicating the version or reward.')

heading('Professional work and daily life')
para('The existing 14-assignment mission catalog uses staged professional workflows. These combine evidence examination, document processing, diagnostic tests, equipment operation, investigation, dialogue, logistics, and creative drafting. A selected procedure is a plan until its required outputs are validated. Some older assignment reports use specified factual terms and length checks; Broad Street uses its separate structured citation review.')
para('Routine employment contracts, special clearance-based assignments, and civic quests remain distinct. Accepted stages produce scoped receipts. Nonclinical routine shifts can repeat once per world cycle with fresh work-order IDs and fresh outputs; one-time work cannot reopen for another payment. Minor errors permit correction, while configured unsafe actions can reduce evaluation and require remediation. These are prototype balancing values rather than final career law.')
para('Tasks, Quests, Missions, Adventures, and Journeys differ in scope without mandatory nesting. Engineering has continuing chapter progress; an Old Factory Ward Adventure and a separate emergent field-response branch provide prototype examples. Existing accomplishments can support several milestones without duplicating their base reward. HOMECRAFT routines and contextual ATLAS discoveries remain possible without turning every action into an assignment.')

heading('Records identity and discovery')
para('AURA is the canonical identity owner. Citizen Hub reads the player\'s explicitly linked identity and eligible personal issuance history. It does not seed a population or adopt whichever person is open in DGAOS. Document Studio retains issuance metadata as historical records rather than rebuilding past issues from today\'s identity.')
para('VAULT distinguishes document types and preserves content and provenance across mail and saved copies. Encrypted records only register inspection when readable. A system driver reference or unrelated archive experience no longer proves lineage or triggers its protected dream. Existing completed discoveries remain in older saves. Document action buttons navigate to owning services instead of awarding repeated credits or silently granting official verification.')
para('TRACE and MEDSTAT await canonical citizen records and listen to the shared update events used by the registry and Document Studio. Selecting and following a subject uses the existing citizen context. These repairs improve continuity between screens; they do not add new institutional powers.')

heading('Presentation and accessibility')
para('The desktop fills the browser without side letterboxing. Normal app windows fit a 16 to 9 frame; maximizing uses the desktop bounds and restoring returns to normal geometry. DGAOS panels respond to their window width. PHONE and FEED use stronger labels, and Directory is a labeled grid button in the toolbar rather than a sidebar entry. Its active state and keyboard focus remain visible.')
para('Appearance offers Mountain Heather, Stillwater Pine, Velvet Sienna, Autumn Study, and Morning Mist, saved with desktop preferences. The default wallpaper is CSS-based and remains available offline; photo wallpapers can keep their image while the surrounding palette changes. Professional screens retain meaningful functional colors. A minimal shared border reset removes accidental browser bevels without removing intentional component borders or focus indicators.')

heading('Saving offline play and recovery')
para('Local browser saves retain gameplay and application records. A loopback .NET service stores an allowlisted world snapshot in SQLite at %LOCALAPPDATA%\\MiraverseOSx\\world-saves.db. Snapshot records cover core progress and reward receipts, housing, assignment documents, browser state, correspondence and phone drafts, canonical identity data, DGAOS and TRACE state, and the creator journal. Individual storage features retain their existing scope; this is not a claim that every file or preference is synchronized.')
para('Offline changes stay on the device and retry when the backend returns. Save and sync displays status, manual saving, and conflict recovery. Revision checks prevent a stale upload from silently replacing a newer world. When both copies have changed, the player chooses which to continue with; displaced device records receive recovery backups and SQLite retains the latest 20 successful revisions. Reconnection retries back off rather than continuously flooding the console.')
para('Broad Street saves acquired IDs, witness request state, formatted draft, structured findings and citations, content version, and submitted reviews with the work order. Closing apps, reloading, or reconnecting does not regenerate authored facts. Unsupported content versions preserve progress but cannot receive automatic review. Local full-access credentials preserve the character, economy, and cases; they do not clear saves or repay completed work.')

doc.add_page_break()
heading('XXV Development Record and Verification',1)
para('The local implementation uses React 19, TypeScript, Vite 6, Tailwind CSS 4, Zustand, and a Photino .NET desktop host. Python builds deterministic authored content during development. It is not a runtime dependency for the investigation. Tiptap supplies local formatted composition without a cloud editor account.')

heading('Content generation and operation')
para('npm start or npm run dev at the repository root starts both the backend on port 5000 and the frontend on port 3000. Starting the frontend alone does not start the save API. npm run desktop starts the native host. npm run build builds the frontend and backend; npm run build:frontend builds the browser assets.')
para('npm run data:case builds Broad Street and npm run test:case validates it. The default seed is 42. A different seed changes an incidental reference while leaving the evidence relationships and solution fixed. Runtime visits never regenerate the package. Its authored assignment is merged by stable ID separately from the database-exported mission catalog, preventing an export from erasing or duplicating it.')
para('npm run data:fiction produces report outlines and separate fictional writing examples; these do not become live AURA citizens, official cases, or paid work. npm run test:fiction checks that generation. npm run data:build remains the existing SQLite data export.')

heading('Latest recorded checks')
para('The October 23 cross-system audit records 177 passing frontend tests, 13 isolated backend integration checks, and four Python tests across case and fictional-content validation. Frontend production and .NET Release builds passed; the Release build reported no warnings or errors. The subsequent shared-border fix also passed the frontend production build. These are recorded results for those revisions, not a claim that this document edit reran all suites.')
para('Broad Street checks cover different source orders, links, eligibility and briefing gates, citation rejection, witness and dispatch deduplication, sanitation and revisions, hydration, version mismatch, legacy completion, and once-only compensation. Its recorded localhost verification completed the five-app workflow. The broader audit combined source review and automated checks rather than manually completing every mission. Large frontend bundle advisories remain.')

heading('Cleanup and compatibility')
para('The unused standalone TRACE wrapper and obsolete folder-structure snapshot were removed. The integrated laboratory, old routes, saved cases, content assets, recovery backups, and standalone service export source were retained. Mail attachment consistency, broken citizen links, premature story discovery, repeat document rewards, and the creator snapshot allowlist mismatch were corrected.')

heading('Open implementation work')
para('Full institution and agency access enforcement, complete student enrollment, richer banking and healthcare systems, deeper household and relationship simulation, further professional content, MAI\'s broader autonomy, Pulse\'s independent purpose, and Creator Studio pause and publish behavior remain open. Browser embedding compatibility is not verified for all eleven external links. Local administrator access is not production authentication or a hosted account service.')
para('The older Phaser Tunnel 09 prototype remains in the codebase. It is not required for Broad Street and does not change the current direction toward desktop investigation. No further Phaser or playable-scene work is included in this update. Courtroom simulation, real internet research, creator publishing, and personal relationship consequences remain outside the Broad Street pilot.')
para('Historical mysteries remain unresolved: White and the Unbound, the Original Creator, MAI\'s complete origins, the final moon names, and detailed House Tribunal procedures. Runtime names, records, numerical tuning, and prototype scenarios do not silently settle those decisions.')

heading('Project references')
para('The following repository records support the implementation chapters and provide locations for future updates. Paths are relative to the MiraverseOSx project folder.')
for text in [
    'docs/application-cleanup-phase-one.md through application-cleanup-phase-three.md: service ownership, Citizen Hub, accounts, memberships, and phone migration.',
    'docs/connected-broad-street.md and docs/broad-street-player-guide.md: case workflow, sources, review, persistence, and compensation.',
    'docs/gameplay-decision-register.md, docs/gameplay-register-implementation.md, and docs/gameplay-018-019-implementation.md: approved structure, professional workflows, daily life, and prototype boundaries.',
    'docs/backend-world.md, docs/appearance.md, docs/tiptap-python.md, and docs/game-audit-2026-10-23.md: saves, design collection, editor and generation tooling, and the recorded audit.',
    'frontend/src/dgaos/components/dga/dga-shell.tsx and DgaServiceWorkspace.tsx: Directory toolbar placement and compact embedded directory.',
    'frontend/src/components/widgets/documentModel.ts, DocumentContent.tsx, and documentAccess.ts: distinct readers, access checks, and owning-service actions.',
    'frontend/src/gameplay/localCredentials.ts, frontend/src/components/ui/windowGeometry.ts, and frontend/src/index.css: local full access, window fit, and shared border reset.',
    'frontend/src/data/broad-street-case.json, broad-street-public.json, and broad-street-assignment.json: versioned authored sources, isolated public publication, and assignment identity.',
]: para(text)

# Keep the reference's typography and margins while fixing inherited title rules,
# inconsistent heading color, duplicate footers, and table pagination.
for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    style=doc.styles[name]
    style.font.color.rgb=RGBColor(0,0,0)
    if style.element.pPr is not None:
        for el in list(style.element.pPr):
            if el.tag==qn('w:pBdr'): style.element.pPr.remove(el)
    if name.startswith('Heading'): style.paragraph_format.keep_with_next=True
doc.styles['Heading 3'].font.name='Aptos'
doc.styles['Heading 3'].font.size=Pt(10)
doc.styles['Heading 3'].font.bold=True

for paragraph in doc.paragraphs:
    is_heading=paragraph.style.name.startswith('Heading')
    if paragraph.style.name=='Subtitle':
        paragraph.paragraph_format.keep_with_next=True
    if is_heading:
        cleaned=re.sub(r'[^\w\s]', ' ', paragraph.text, flags=re.UNICODE)
        cleaned=re.sub(r'\s+', ' ', cleaned).strip()
        if cleaned!=paragraph.text: paragraph.text=cleaned
    if is_heading or paragraph.style.name in ('Title','Subtitle'):
        for run in paragraph.runs: run.font.color.rgb=RGBColor(0,0,0)
        pr=paragraph._p.pPr
        if pr is not None:
            for el in list(pr):
                if el.tag==qn('w:pBdr'): pr.remove(el)
    paragraph.paragraph_format.widow_control=True

# Remove inherited decorative paragraph borders from the cover block.
for paragraph in doc.paragraphs[:5]:
    if paragraph._p.pPr is not None:
        for el in list(paragraph._p.pPr):
            if el.tag==qn('w:pBdr'): paragraph._p.pPr.remove(el)

for t in doc.tables:
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    old=t._tbl.tblPr.find(qn('w:tblBorders'))
    if old is not None:t._tbl.tblPr.remove(old)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr()
        if ri==0:
            mark=OxmlElement('w:tblHeader');pr.append(mark)
        cant=OxmlElement('w:cantSplit');pr.append(cant)
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            old=tcpr.find(qn('w:tcBorders'))
            if old is not None:tcpr.remove(old)
            cb=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');cb.append(b)
            tcpr.append(cb)
            old=tcpr.find(qn('w:shd'))
            if old is not None:tcpr.remove(old)
            shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E8ECF0' if ri==0 else 'FFFFFF');tcpr.append(shd)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before=Pt(4)
                paragraph.paragraph_format.space_after=Pt(4)
                paragraph.paragraph_format.keep_with_next=(ri==0)
                for run in paragraph.runs:
                    run.font.color.rgb=RGBColor(0,0,0)
                    if ri==0:run.bold=True

for section in doc.sections:
    footer=section.footer
    for el in list(footer._element):footer._element.remove(el)
    paragraph=footer.add_paragraph()
    paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=paragraph.add_run('MERIDION  |  WORLD AND APPLICATION BIBLE  |  ')
    r.font.name='Aptos';r.font.size=Pt(8);r.font.color.rgb=RGBColor(0,0,0)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');paragraph._p.append(field)

doc.core_properties.title='Meridion World Lore and Application Bible'
doc.core_properties.subject='World lore and connected desktop implementation as of October 23 2026'
doc.core_properties.comments=''
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT)
print('Updated existing paragraphs:',len(replacements),'Total paragraphs:',len(doc.paragraphs),'Tables:',len(doc.tables))

# Confirm original historical substance survives and revised implementation exists.
check=Document(OUT)
assert check.paragraphs[23+13].text==p[23].text
text='\n'.join(x.text for x in check.paragraphs)
for required in ['ms-case-002','clearance 2','10-stamina','Phone Messages','177 passing','not production authentication']:
    assert required in text,required
assert 'relationship to other identity interfaces remains an area' not in text
print('Content checks passed')
